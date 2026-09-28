"""
Builds docs/NeuroLens_Pitch_Deck.pptx (10 slides) with python-pptx.

The prototype-output slide runs the real rule-based compiler, so the slide
always matches what `python cognitive_compiler.py --no-model` prints.

Usage:  python generate_deck.py
"""
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from cognitive_compiler import CognitiveCompiler

# ---- palette: deep navy dominant, teal = implemented, amber = planned/inferred ----
NAVY = RGBColor(0x0B, 0x1F, 0x33)
INK = RGBColor(0x1B, 0x2A, 0x3A)
SLATE = RGBColor(0x55, 0x65, 0x75)
TEAL = RGBColor(0x0E, 0x8F, 0x83)
TEAL_TINT = RGBColor(0xDD, 0xF1, 0xEE)
AMBER = RGBColor(0xF2, 0xB1, 0x34)
ROSE = RGBColor(0xC0, 0x4B, 0x3B)
PAPER = RGBColor(0xF5, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MIST = RGBColor(0xC9, 0xD6, 0xE3)
LINE = RGBColor(0xB8, 0xC4, 0xD0)

HEAD_FONT = "Cambria"
BODY_FONT = "Calibri"

MEETING = (
    "It would probably be good if you could get the analysis over to me sometime before Friday. "
    "And maybe revisit the first section because I don't think we're quite there yet."
)


# ------------------------------------------------------------------ helpers ----
def bg(slide, color):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def box(slide, x, y, w, h, fill=None, line=None, dash=False, rounded=True, line_w=1.25):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    if rounded:
        shape.adjustments[0] = 0.06
    shape.shadow.inherit = False
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
        if dash:
            shape.line.dash_style = MSO_LINE.DASH
    return shape


def write(shape_or_frame, paras, anchor=MSO_ANCHOR.TOP, margins=(0.15, 0.1, 0.15, 0.1), align=PP_ALIGN.LEFT):
    """paras: list of paragraphs; each paragraph is a list of (text, size, bold, color[, italic, font])."""
    tf = shape_or_frame.text_frame if hasattr(shape_or_frame, "text_frame") else shape_or_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left, tf.margin_top, tf.margin_right, tf.margin_bottom = [Inches(m) for m in margins]
    for i, segs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i > 0:
            p.space_before = Pt(6)
        for seg in segs:
            text, size, bold, color = seg[:4]
            r = p.add_run()
            r.text = text
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
            r.font.italic = seg[4] if len(seg) > 4 else False
            r.font.name = seg[5] if len(seg) > 5 else BODY_FONT
    return tf


def label(slide, x, y, w, h, paras, **kw):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    write(tb, paras, margins=kw.pop("margins", (0, 0, 0, 0)), **kw)
    return tb


def title(slide, text, dark=False):
    color = WHITE if dark else NAVY
    label(slide, 0.7, 0.5, 11.9, 0.9, [[(text, 34, True, color, False, HEAD_FONT)]], anchor=MSO_ANCHOR.MIDDLE)


def arrow(slide, x1, y1, x2, y2, color=SLATE, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    tail = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
    tail.set("type", "triangle")
    return c


def badge(slide, x, y, d, text, fill, color=WHITE, size=16):
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.shadow.inherit = False
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    write(c, [[(text, size, True, color)]], anchor=MSO_ANCHOR.MIDDLE, margins=(0, 0, 0, 0), align=PP_ALIGN.CENTER)
    return c


def node(slide, x, y, w, h, head, sub, status):
    """Architecture node. status: 'done' (teal, solid) or 'planned' (white, dashed)."""
    if status == "done":
        s = box(slide, x, y, w, h, fill=TEAL)
        hc, sc = WHITE, TEAL_TINT
    else:
        s = box(slide, x, y, w, h, fill=WHITE, line=SLATE, dash=True)
        hc, sc = INK, SLATE
    paras = [[(head, 16, True, hc)]]
    if sub:
        paras.append([(sub, 12, False, sc)])
    write(s, paras, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
    return s


# -------------------------------------------------------------------- deck ----
def build(output_path="docs/NeuroLens_Pitch_Deck.pptx"):
    insight = CognitiveCompiler(use_model=False).compile(MEETING)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1 ── Title ───────────────────────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    label(s, 0.9, 2.0, 11.5, 1.4, [[("NeuroLens", 66, True, WHITE, False, HEAD_FONT)]])
    label(s, 0.9, 3.4, 11.5, 0.7, [[("Private, real-time AI for cognitive accessibility", 26, False, AMBER)]])
    label(
        s, 0.9, 4.3, 11.5, 0.6,
        [[("Understand the moment. Organize the thought. Keep everything private.", 18, False, MIST, True)]],
    )
    label(
        s, 0.9, 6.2, 11.5, 0.6,
        [[("Snapdragon® AI Lab Build & Present Challenge 2026   |   Swati Dubey   |   Proposal + tested prototype", 14, False, MIST)]],
    )

    # 2 ── Problem ─────────────────────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s, PAPER)
    title(s, "Fast, indirect communication hides the real task")
    card = box(s, 0.7, 1.9, 6.0, 3.9, fill=WHITE, line=LINE)
    write(
        card,
        [
            [("What was said", 13, True, SLATE)],
            [("“It would probably be good if you could get the analysis over to me sometime before Friday.”", 20, False, INK, True, HEAD_FONT)],
            [("", 8, False, INK)],
            [("What a person has to work out, live", 13, True, SLATE)],
            [("Is this a request or a suggestion?", 17, False, INK)],
            [("What is the deadline, exactly?", 17, False, INK)],
            [("What else is unclear?", 17, False, INK)],
        ],
        margins=(0.35, 0.3, 0.35, 0.3),
    )
    rows = [
        ("1", "Indirect wording", "Polite phrasing can hide real requests and deadlines."),
        ("2", "No time to parse", "Lectures and meetings move on before you have worked it out."),
        ("3", "Cloud help costs privacy", "Assistants that fix this usually need your conversations sent to a remote server."),
    ]
    y = 1.9
    for n, head, body in rows:
        badge(s, 7.2, y + 0.1, 0.6, n, TEAL)
        label(s, 8.05, y, 4.6, 1.4, [[(head, 20, True, NAVY)], [(body, 15, False, SLATE)]])
        y += 1.6

    # 3 ── Gap ─────────────────────────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s, PAPER)
    title(s, "What is missing today")
    cards = [
        ("Transcription only", "Gives you every word.", "Does not say what matters or what to do."),
        ("Cloud AI assistants", "Can summarize and answer.", "Needs private conversations sent off-device."),
        ("Task apps", "Organize tasks well.", "Start from tasks you type in yourself."),
        ("Overconfident answers", "Sound certain.", "Rarely say what is unclear or ask instead."),
    ]
    x = 0.7
    for head, does, gap in cards:
        c = box(s, x, 1.9, 2.85, 3.5, fill=WHITE, line=LINE)
        write(
            c,
            [
                [(head, 19, True, NAVY, False, HEAD_FONT)],
                [("", 6, False, INK)],
                [(does, 15, False, SLATE)],
                [("", 6, False, INK)],
                [("Gap", 12, True, ROSE)],
                [(gap, 16, False, INK)],
            ],
            margins=(0.25, 0.3, 0.25, 0.25),
        )
        x += 3.05

    # 4 ── Solution flow ───────────────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s, PAPER)
    title(s, "From messy language to structure")
    steps = [
        ("Parse", "Split speech or text into clauses. Find dates and times with rules."),
        ("Disambiguate", "Score how vague the wording is. Spot missing task, deadline or referent."),
        ("Structure", "Fill a fixed schema: actions, deadlines, interpretations, questions."),
        ("Validate", "Output must pass the schema, or the system falls back to rules."),
    ]
    x = 0.7
    for i, (head, body) in enumerate(steps):
        n = box(s, x, 2.0, 2.7, 0.9, fill=NAVY)
        write(n, [[(f"{i + 1}  {head}", 20, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER)
        label(s, x, 3.15, 2.7, 2.0, [[(body, 16, False, INK)]])
        if i < 3:
            arrow(s, x + 2.75, 2.45, x + 3.05, 2.45, color=SLATE)
        x += 3.05
    note = box(s, 0.7, 5.5, 11.9, 1.1, fill=TEAL_TINT)
    write(
        note,
        [[("Not a chatbot. ", 18, True, NAVY), ("It returns the same fixed structure every time, so every field can be shown, checked and reused.", 18, False, INK)]],
        anchor=MSO_ANCHOR.MIDDLE,
        margins=(0.35, 0.1, 0.35, 0.1),
    )

    # 5 ── Known / Inferred / Unknown ─────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s, PAPER)
    title(s, "Known  ≠  Inferred  ≠  Unknown")
    kiu = [
        ("KNOWN", TEAL, WHITE, "What is stated", "Literal meaning, explicit tasks, deadlines found by rules.", "Deadline: Friday"),
        ("INFERRED", AMBER, INK, "What is possible", "Readings of vague wording, always with the words that support them.", "May be a real request, not just a suggestion"),
        ("UNKNOWN", ROSE, WHITE, "What is missing", "Gaps become a question, not a guess.", "“Which part should be changed, and how?”"),
    ]
    x = 0.7
    for head, fill, fg, sub, body, example in kiu:
        c = box(s, x, 1.9, 3.85, 4.5, fill=fill)
        write(
            c,
            [
                [(head, 30, True, fg, False, HEAD_FONT)],
                [(sub, 16, True, fg)],
                [("", 6, False, fg)],
                [(body, 17, False, fg)],
                [("", 10, False, fg)],
                [("From the meeting example", 12, True, fg)],
                [(example, 16, False, fg, True)],
            ],
            margins=(0.35, 0.35, 0.35, 0.3),
        )
        x += 4.02

    # 6 ── Architecture ────────────────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s, WHITE)
    title(s, "Architecture: what exists and what is planned")
    node(s, 0.7, 1.7, 1.9, 1.0, "Microphone", None, "planned")
    node(s, 3.0, 1.7, 1.6, 1.0, "VAD", None, "planned")
    node(s, 5.0, 1.7, 2.0, 1.0, "Whisper-Base", "speech to text", "planned")
    node(s, 0.7, 3.3, 1.9, 1.0, "Screenshot", "user-selected", "planned")
    node(s, 5.0, 3.3, 2.0, 1.0, "Qwen3-VL-4B", "screen to text", "planned")
    node(s, 5.0, 4.9, 2.0, 1.0, "Typed text", None, "done")
    node(s, 7.7, 2.8, 2.5, 2.2, "Cognitive Compiler", "rules + local model (GenieX)", "done")
    node(s, 10.7, 3.4, 2.0, 1.0, "Schema validation", "Pydantic", "done")
    node(s, 10.7, 5.0, 2.0, 1.1, "Tasks, Focus, memory, UI", None, "planned")
    arrow(s, 2.6, 2.2, 3.0, 2.2)
    arrow(s, 4.6, 2.2, 5.0, 2.2)
    arrow(s, 2.6, 3.8, 5.0, 3.8)
    arrow(s, 7.0, 2.2, 7.7, 3.4)
    arrow(s, 7.0, 3.8, 7.7, 3.9)
    arrow(s, 7.0, 5.4, 7.7, 4.5)
    arrow(s, 10.2, 3.9, 10.7, 3.9)
    arrow(s, 11.7, 4.4, 11.7, 5.0)
    key1 = box(s, 0.7, 6.55, 0.3, 0.3, fill=TEAL, rounded=False)
    label(s, 1.1, 6.5, 3.6, 0.4, [[("Implemented and tested on CPU", 13, False, INK)]], anchor=MSO_ANCHOR.MIDDLE)
    key2 = box(s, 4.8, 6.55, 0.3, 0.3, fill=WHITE, line=SLATE, dash=True, rounded=False)
    label(s, 5.2, 6.5, 2.0, 0.4, [[("Planned", 13, False, INK)]], anchor=MSO_ANCHOR.MIDDLE)
    label(
        s, 7.0, 6.5, 5.7, 0.4,
        [[("Model path written; not yet run on my Snapdragon device", 13, False, SLATE, True)]],
        anchor=MSO_ANCHOR.MIDDLE,
    )

    # 7 ── Snapdragon plan ─────────────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s, PAPER)
    title(s, "Snapdragon plan: targets, not results")
    rows = [
        ("Role", "Candidate model", "Intended runtime", "Status"),
        ("Speech to text", "Whisper-Base", "Qualcomm AI Hub / QAI AppBuilder", "Not yet run on my device"),
        ("Reasoning", "Qwen3-0.6B (Phi-4-Mini alt.)", "GenieX local server", "Client written; not yet run"),
        ("Screen understanding", "Qwen3-VL-4B-Instruct", "GenieX", "Planned"),
    ]
    tbl = s.shapes.add_table(4, 4, Inches(0.7), Inches(1.9), Inches(11.9), Inches(3.0)).table
    for ci, w in enumerate([2.6, 3.2, 3.4, 2.7]):
        tbl.columns[ci].width = Inches(w)
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = Inches(0.75)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY if ri == 0 else (WHITE if ri % 2 else TEAL_TINT)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.15)
            run = tf.paragraphs[0].add_run()
            run.text = val
            run.font.name = BODY_FONT
            run.font.size = Pt(16)
            run.font.bold = ri == 0 or ci == 0
            run.font.color.rgb = WHITE if ri == 0 else INK
    call = box(s, 0.7, 5.35, 11.9, 1.2, fill=WHITE, line=LINE)
    write(
        call,
        [[("Inference is intended to run on the Snapdragon NPU where supported; UI, storage and rules run on the CPU. ", 16, False, INK),
          ("No speed or accuracy numbers are claimed until measured on my own laptop.", 16, True, NAVY)]],
        anchor=MSO_ANCHOR.MIDDLE,
        margins=(0.35, 0.1, 0.35, 0.1),
    )

    # 8 ── Privacy ─────────────────────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s, PAPER)
    title(s, "Privacy by design (goals for the full app)")
    items = [
        ("Ephemeral by default", "Audio and screenshots are processed in memory and discarded unless the user chooses to save."),
        ("Preferences, not logs", "Memory keeps user-approved settings such as “prefers short answers”, not conversation history."),
        ("No cloud AI for core use", "Local models handle inference. Offline behaviour will be tested before it is claimed."),
    ]
    y = 1.9
    for i, (head, body) in enumerate(items):
        c = box(s, 0.7, y, 11.9, 1.3, fill=WHITE, line=LINE)
        badge(s, 1.0, y + 0.3, 0.7, str(i + 1), NAVY, size=20)
        label(s, 2.0, y + 0.15, 10.3, 1.1, [[(head, 20, True, NAVY)], [(body, 16, False, SLATE)]], anchor=MSO_ANCHOR.MIDDLE)
        y += 1.5
    label(s, 0.7, 6.5, 11.9, 0.5, [[("These are design goals; the privacy features are not built yet.", 14, False, SLATE, True)]])

    # 9 ── Prototype output (real) ─────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s, PAPER)
    title(s, "Prototype today: real output from the rule-based engine")
    left = box(s, 0.7, 1.8, 4.9, 3.9, fill=WHITE, line=LINE)
    write(
        left,
        [[("Input", 13, True, SLATE)],
         [("“" + MEETING + "”", 16, False, INK, True, HEAD_FONT)]],
        margins=(0.3, 0.3, 0.3, 0.3),
    )
    deadline = ", ".join(sorted({d.deadline for d in insight.deadlines})) or "none"
    right = box(s, 5.9, 1.8, 6.7, 3.9, fill=WHITE, line=TEAL, line_w=2)
    paras = [[("Output", 13, True, TEAL)], [("Actions", 13, True, SLATE)]]
    paras += [[(f"[ ]  {a.task}  ({a.priority})", 16, False, INK)] for a in insight.actions]
    paras += [
        [("Deadline", 13, True, SLATE)],
        [(deadline, 16, False, INK)],
        [("Ambiguity", 13, True, SLATE)],
        [(insight.ambiguity.level.upper(), 16, True, ROSE)],
        [("Clarifying question", 13, True, SLATE)],
        [(f"“{insight.clarifying_question}”", 16, False, INK, True)],
    ]
    write(right, paras, margins=(0.35, 0.25, 0.35, 0.2))
    label(
        s, 0.7, 6.0, 11.9, 0.6,
        [[("Produced on CPU by the rule engine (python cognitive_compiler.py --no-model). The local-model path has not yet been run on Snapdragon.", 13, False, SLATE, True)]],
    )

    # 10 ── Status and next ────────────────────────────────────────────────────
    s = prs.slides.add_slide(blank)
    bg(s, NAVY)
    title(s, "Where it stands and what comes next", dark=True)
    cols = [
        ("Done", TEAL, WHITE,
         ["Structured-output schema", "Rule-based compiler", "Model client with validation and fallback", "16 unit tests"]),
        ("Next", AMBER, INK,
         ["Run GenieX + Qwen3 on my laptop", "Whisper-Base on audio", "Task creation, Focus Mode, simple UI", "Screenshot input with Qwen3-VL"]),
        ("To measure", WHITE, INK,
         ["Task and deadline accuracy on ~30 labelled samples", "Latency: ASR, first token, end to end", "CPU vs NPU", "Offline test"]),
    ]
    x = 0.7
    for head, fill, fg, lines in cols:
        c = box(s, x, 1.9, 3.85, 3.9, fill=fill)
        paras = [[(head, 26, True, fg, False, HEAD_FONT)]] + [[("•  " + ln, 16, False, fg)] for ln in lines]
        write(c, paras, margins=(0.35, 0.3, 0.3, 0.3))
        x += 4.02

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    prs.save(output_path)
    print(f"Saved {output_path}")


if __name__ == "__main__":
    build()
